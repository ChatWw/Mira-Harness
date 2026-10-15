"""Mira adapter for XlsxWriter (BSD-2-Clause) and openpyxl (MIT)."""
import json
import sys
from pathlib import Path


def main():
    action, source = sys.argv[1:3]
    if action == "inspect":
        from openpyxl import load_workbook
        workbook = load_workbook(source, read_only=True, data_only=False)
        try:
            print(json.dumps({"sheets": [{"name": sheet.title, "rows": list(sheet.values)} for sheet in workbook]}, ensure_ascii=False, default=str))
        finally:
            workbook.close()
    elif action == "create":
        import xlsxwriter
        target = Path(sys.argv[3])
        if target.exists():
            raise FileExistsError(f"Refusing to overwrite {target}")
        data = json.loads(Path(source).read_text(encoding="utf-8-sig"))
        target.parent.mkdir(parents=True, exist_ok=True)
        with xlsxwriter.Workbook(target, {"strings_to_formulas": False, "strings_to_urls": False}) as workbook:
            header = workbook.add_format({"bold": True, "bg_color": "#E8EDF2", "bottom": 1})
            for item in data["sheets"]:
                sheet = workbook.add_worksheet(item["name"])
                sheet.freeze_panes(1, 0)
                sheet.set_column(0, max((len(row) for row in item["rows"]), default=1) - 1, 18)
                for row_index, row in enumerate(item["rows"]):
                    for column_index, value in enumerate(row):
                        if isinstance(value, dict) and "formula" in value:
                            if "value" not in value:
                                raise ValueError("Formula cells require an explicitly verified cached value")
                            sheet.write_formula(row_index, column_index, value["formula"], None, value["value"])
                        else:
                            sheet.write(row_index, column_index, value, header if row_index == 0 else None)
        print(str(target.resolve()))
    else:
        raise ValueError("Expected create or inspect")


if __name__ == "__main__":
    main()
